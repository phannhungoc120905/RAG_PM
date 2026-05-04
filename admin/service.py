import hashlib
import json
import secrets
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException, status
from sqlalchemy import Select, delete, func, or_, select
from sqlalchemy.orm import Session, selectinload

from auth.service import hash_password
from config import settings
from db.models import (
    APIKey,
    ChunkMetadata,
    Document,
    PermissionGroup,
    PermissionGroupFunction,
    Department,
    Document,
    IssuingUnit,
    LoginHistory,
    NoticeDocument,
    PermissionGroup,
    PermissionGroupFunction,
    Position,
    SummaryHistory,
    SystemConfig,
    SystemFunction,
    SystemLog,
    User,
    WorkAssignmentDocument,
    WorkItem,
)


SNAPSHOT_MODELS = [
    PermissionGroup,
    SystemFunction,
    PermissionGroupFunction,
    User,
    IssuingUnit,
    Department,
    Position,
    User,
    LoginHistory,
    SystemConfig,
    Document,
    ChunkMetadata,
    SummaryHistory,
    WorkAssignmentDocument,
    WorkItem,
    NoticeDocument,
    SystemLog,
    APIKey,
]

SNAPSHOT_DELETE_ORDER = [
    PermissionGroupFunction,
    APIKey,
    SystemLog,
    SummaryHistory,
    ChunkMetadata,
    Document,
    SystemConfig,
    User,
    WorkItem,
    NoticeDocument,
    WorkAssignmentDocument,
    SummaryHistory,
    ChunkMetadata,
    Document,
    LoginHistory,
    SystemConfig,
    User,
    Position,
    Department,
    IssuingUnit,
    SystemFunction,
    PermissionGroup,
]


def list_users(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
) -> tuple[list[User], int]:
    statement: Select[tuple[User]] = select(User).options(
        selectinload(User.permission_group),
        selectinload(User.department),
        selectinload(User.position),
    )
    count_statement = select(func.count()).select_from(User)

    if search:
        criteria = or_(
            User.username.ilike(f"%{search}%"),
            User.email.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)

    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(User.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_user(db: Session, payload: dict[str, Any]) -> User:
    _assert_unique_user(db, payload["username"], payload.get("email"))
    permission_group_id = payload.get("permission_group_id")
    if permission_group_id is not None:
        _get_group_or_404(db, permission_group_id)
    _validate_user_relations(
        db,
        payload.get("permission_group_id"),
        payload.get("department_id"),
        payload.get("position_id"),
    )
    user = User(
        username=payload["username"],
        email=payload.get("email"),
        hashed_password=hash_password(payload["password"]),
        role=payload.get("role", "user"),
        permission_group_id=payload.get("permission_group_id"),
        department_id=payload.get("department_id"),
        position_id=payload.get("position_id"),
        is_active=payload.get("is_active", True),
        auth_source=payload.get("auth_source", "local"),
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def update_user(db: Session, user_id: int, payload: dict[str, Any]) -> User:
    user = _get_user_or_404(db, user_id)
    username = payload.get("username")
    email = payload.get("email")

    if username and username != user.username:
        if db.scalar(select(User).where(User.username == username)):
            raise HTTPException(status_code=409, detail="Username already exists")
        user.username = username

    if email != user.email:
        if email and db.scalar(select(User).where(User.email == email, User.id != user.id)):
            raise HTTPException(status_code=409, detail="Email already exists")
        user.email = email

    for field in ("role", "is_active", "permission_group_id"):
        if field in payload:
            if field == "permission_group_id" and payload[field] is not None:
                _get_group_or_404(db, payload[field])
    _validate_user_relations(
        db,
        payload.get("permission_group_id", user.permission_group_id) if "permission_group_id" in payload else user.permission_group_id,
        payload.get("department_id", user.department_id) if "department_id" in payload else user.department_id,
        payload.get("position_id", user.position_id) if "position_id" in payload else user.position_id,
    )

    for field in ("role", "is_active", "permission_group_id", "department_id", "position_id"):
        if field in payload:
            setattr(user, field, payload[field])

    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def toggle_user_active(
    db: Session,
    target_user_id: int,
    actor_user_id: int,
) -> User:
    if target_user_id == actor_user_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="You cannot deactivate your own account",
        )

    user = _get_user_or_404(db, target_user_id)
    user.is_active = not user.is_active
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def reset_user_password(db: Session, user_id: int, new_password: str) -> None:
    user = _get_user_or_404(db, user_id)
    user.hashed_password = hash_password(new_password)
    db.add(user)
    db.commit()


def list_logs(
    db: Session,
    page: int,
    page_size: int,
    user_id: int | None,
    action: str | None,
    date_from: datetime | None,
    date_to: datetime | None,
) -> tuple[list[SystemLog], int]:
    statement: Select[tuple[SystemLog]] = select(SystemLog)
    count_statement = select(func.count()).select_from(SystemLog)

    if user_id is not None:
        statement = statement.where(SystemLog.user_id == user_id)
        count_statement = count_statement.where(SystemLog.user_id == user_id)
    if action:
        statement = statement.where(SystemLog.action.ilike(f"%{action}%"))
        count_statement = count_statement.where(SystemLog.action.ilike(f"%{action}%"))
    if date_from:
        statement = statement.where(SystemLog.created_at >= date_from)
        count_statement = count_statement.where(SystemLog.created_at >= date_from)
    if date_to:
        statement = statement.where(SystemLog.created_at <= date_to)
        count_statement = count_statement.where(SystemLog.created_at <= date_to)

    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(SystemLog.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all())
    return items, total


def list_login_history(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    user_id: int | None,
    login_type: str | None,
    status_value: str | None,
    date_from: datetime | None,
    date_to: datetime | None,
) -> tuple[list[LoginHistory], int]:
    statement: Select[tuple[LoginHistory]] = select(LoginHistory).options(selectinload(LoginHistory.user))
    count_statement = select(func.count()).select_from(LoginHistory)

    if search:
        criteria = LoginHistory.username_snapshot.ilike(f"%{search}%")
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if user_id is not None:
        statement = statement.where(LoginHistory.user_id == user_id)
        count_statement = count_statement.where(LoginHistory.user_id == user_id)
    if login_type:
        statement = statement.where(LoginHistory.login_type == login_type)
        count_statement = count_statement.where(LoginHistory.login_type == login_type)
    if status_value:
        statement = statement.where(LoginHistory.status == status_value)
        count_statement = count_statement.where(LoginHistory.status == status_value)
    if date_from:
        statement = statement.where(LoginHistory.login_at >= date_from)
        count_statement = count_statement.where(LoginHistory.login_at >= date_from)
    if date_to:
        statement = statement.where(LoginHistory.login_at <= date_to)
        count_statement = count_statement.where(LoginHistory.login_at <= date_to)

    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(LoginHistory.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all())
    return items, total


def export_logs_to_excel(
    db: Session,
    user_id: int | None,
    action: str | None,
    date_from: datetime | None,
    date_to: datetime | None,
) -> Path:
    workbook_cls = _get_workbook_class()
    logs, _ = list_logs(db, 1, 100000, user_id, action, date_from, date_to)

    export_dir = Path(settings.BACKUP_DIR) / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    export_path = export_dir / f"system_logs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    workbook = workbook_cls()
    sheet = workbook.active or workbook.create_sheet("system_logs", 0)
    sheet.title = "system_logs"
    sheet.append(["id", "user_id", "action", "detail", "ip_address", "status_code", "created_at"])
    for item in logs:
        sheet.append(
            [
                item.id,
                item.user_id,
                item.action,
                item.detail,
                item.ip_address,
                item.status_code,
                item.created_at.isoformat() if item.created_at else None,
            ]
        )
    workbook.save(export_path)
    return export_path


def get_system_status(db: Session) -> dict[str, Any]:
    upload_dir = Path(settings.UPLOAD_DIR)
    disk_target = upload_dir if upload_dir.exists() else Path(".")
    usage = shutil.disk_usage(disk_target)

    return {
        "disk": {
            "total_gb": round(usage.total / (1024**3), 2),
            "used_gb": round(usage.used / (1024**3), 2),
            "free_gb": round(usage.free / (1024**3), 2),
        },
        "uploads": {
            "processing_count": _count_files(upload_dir / "processing"),
            "done_count": _count_files(upload_dir / "done"),
            "failed_count": _count_files(upload_dir / "failed"),
        },
        "db": {
            "users_count": db.scalar(select(func.count()).select_from(User)) or 0,
            "documents_count": db.scalar(select(func.count()).select_from(Document)) or 0,
            "summaries_count": db.scalar(select(func.count()).select_from(SummaryHistory)) or 0,
        },
        "app": {
            "version": "1.0.0",
            "debug": settings.DEBUG,
        },
    }


def list_system_configs(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    category: str | None,
) -> tuple[list[SystemConfig], int]:
    statement: Select[tuple[SystemConfig]] = select(SystemConfig)
    count_statement = select(func.count()).select_from(SystemConfig)

    if search:
        criteria = or_(
            SystemConfig.config_key.ilike(f"%{search}%"),
            SystemConfig.description.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if category:
        statement = statement.where(SystemConfig.category == category)
        count_statement = count_statement.where(SystemConfig.category == category)

    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(SystemConfig.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all())
    return items, total


def create_system_config(
    db: Session,
    actor_user_id: int,
    payload: dict[str, Any],
) -> SystemConfig:
    if db.scalar(select(SystemConfig).where(SystemConfig.config_key == payload["config_key"])):
        raise HTTPException(status_code=409, detail="Config key already exists")

    config = SystemConfig(
        config_key=payload["config_key"],
        config_value=str(payload["config_value"]),
        category=payload.get("category", "general"),
        data_type=payload.get("data_type", "string"),
        description=payload.get("description"),
        is_active=payload.get("is_active", True),
        updated_by=actor_user_id,
    )
    db.add(config)
    db.commit()
    db.refresh(config)
    _sync_config_to_settings(config.config_key, config.config_value, config.data_type)
    return config


def update_system_config(
    db: Session,
    config_id: int,
    actor_user_id: int,
    payload: dict[str, Any],
) -> SystemConfig:
    config = _get_config_or_404(db, config_id)
    if "config_key" in payload and payload["config_key"] != config.config_key:
        if db.scalar(select(SystemConfig).where(SystemConfig.config_key == payload["config_key"])):
            raise HTTPException(status_code=409, detail="Config key already exists")
    for field in ("config_key", "config_value", "category", "data_type", "description", "is_active"):
        if field in payload:
            setattr(config, field, str(payload[field]) if field == "config_value" else payload[field])
    config.updated_by = actor_user_id
    config.updated_at = datetime.utcnow()
    db.add(config)
    db.commit()
    db.refresh(config)
    _sync_config_to_settings(config.config_key, config.config_value, config.data_type)
    return config


def delete_system_config(db: Session, config_id: int) -> None:
    config = _get_config_or_404(db, config_id)
    db.delete(config)
    db.commit()


def get_safe_config() -> dict[str, Any]:
    return {
        "model_name": settings.MODEL_NAME,
        "model_path": settings.MODEL_PATH,
        "faiss_index_path": settings.FAISS_INDEX_PATH,
        "vector_dim": settings.VECTOR_DIM,
        "chunk_size": settings.CHUNK_SIZE,
        "chunk_overlap": settings.CHUNK_OVERLAP,
        "ocr_lang": settings.OCR_LANG,
    }


def update_runtime_config(payload: dict[str, Any]) -> dict[str, Any]:
    for field in ("model_name", "chunk_size", "chunk_overlap"):
        if field in payload:
            setattr(settings, field.upper(), payload[field])
    return get_safe_config()


def list_api_keys(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
) -> tuple[list[APIKey], int]:
    statement: Select[tuple[APIKey]] = select(APIKey).options(selectinload(APIKey.creator))
    count_statement = select(func.count()).select_from(APIKey)

    if search:
        criteria = or_(
            APIKey.name.ilike(f"%{search}%"),
            APIKey.description.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)

    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(APIKey.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_api_key(
    db: Session,
    actor_user_id: int,
    payload: dict[str, Any],
) -> tuple[APIKey, str]:
    raw_key = secrets.token_urlsafe(32)
    key_hash = hashlib.sha256(raw_key.encode("utf-8")).hexdigest()
    api_key = APIKey(
        name=payload["name"],
        description=payload.get("description"),
        key_hash=key_hash,
        created_by=actor_user_id,
        is_active=payload.get("is_active", True),
    )
    db.add(api_key)
    db.commit()
    db.refresh(api_key)
    return api_key, raw_key


def update_api_key(db: Session, api_key_id: int, payload: dict[str, Any]) -> APIKey:
    api_key = _get_api_key_or_404(db, api_key_id)
    for field in ("name", "description", "is_active"):
        if field in payload:
            setattr(api_key, field, payload[field])
    api_key.updated_at = datetime.utcnow()
    db.add(api_key)
    db.commit()
    db.refresh(api_key)
    return api_key


def delete_api_key(db: Session, api_key_id: int) -> None:
    api_key = _get_api_key_or_404(db, api_key_id)
    db.delete(api_key)
    db.commit()


def list_permission_groups(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
) -> tuple[list[PermissionGroup], int]:
    statement: Select[tuple[PermissionGroup]] = select(PermissionGroup).options(
        selectinload(PermissionGroup.function_permissions).selectinload(
            PermissionGroupFunction.function
        )
    )
    count_statement = select(func.count()).select_from(PermissionGroup)

    if search:
        criteria = or_(
            PermissionGroup.name.ilike(f"%{search}%"),
            PermissionGroup.code.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)

    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(PermissionGroup.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all())
    return items, total


def create_permission_group(db: Session, payload: dict[str, Any]) -> PermissionGroup:
    if db.scalar(select(PermissionGroup).where(PermissionGroup.code == payload["code"])):
        raise HTTPException(status_code=409, detail="Permission group code already exists")

    group = PermissionGroup(
        name=payload["name"],
        code=payload["code"],
        description=payload.get("description"),
        is_active=payload.get("is_active", True),
    )
    db.add(group)
    db.flush()
    _replace_group_permissions(db, group, payload.get("permissions", []))
    db.commit()
    db.refresh(group)
    return group


def update_permission_group(
    db: Session,
    group_id: int,
    payload: dict[str, Any],
) -> PermissionGroup:
    group = _get_group_or_404(db, group_id)
    if "code" in payload and payload["code"] != group.code:
        if db.scalar(select(PermissionGroup).where(PermissionGroup.code == payload["code"])):
            raise HTTPException(status_code=409, detail="Permission group code already exists")
    for field in ("name", "code", "description", "is_active"):
        if field in payload:
            setattr(group, field, payload[field])
    if "permissions" in payload:
        _replace_group_permissions(db, group, payload["permissions"])
    group.updated_at = datetime.utcnow()
    db.add(group)
    db.commit()
    db.refresh(group)
    return group


def delete_permission_group(db: Session, group_id: int) -> None:
    group = _get_group_or_404(db, group_id)
    if group.users:
        raise HTTPException(status_code=400, detail="Group is assigned to users")
    db.delete(group)
    db.commit()


def list_system_functions(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
) -> tuple[list[SystemFunction], int]:
    statement: Select[tuple[SystemFunction]] = select(SystemFunction)
    count_statement = select(func.count()).select_from(SystemFunction)

    if search:
        criteria = or_(
            SystemFunction.name.ilike(f"%{search}%"),
            SystemFunction.code.ilike(f"%{search}%"),
            SystemFunction.module.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)

    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(SystemFunction.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all())
    return items, total


def create_system_function(db: Session, payload: dict[str, Any]) -> SystemFunction:
    if db.scalar(select(SystemFunction).where(SystemFunction.code == payload["code"])):
        raise HTTPException(status_code=409, detail="Function code already exists")

    function = SystemFunction(
        name=payload["name"],
        code=payload["code"],
        module=payload.get("module"),
        description=payload.get("description"),
        is_active=payload.get("is_active", True),
    )
    db.add(function)
    db.commit()
    db.refresh(function)
    return function


def update_system_function(
    db: Session,
    function_id: int,
    payload: dict[str, Any],
) -> SystemFunction:
    function = _get_function_or_404(db, function_id)
    if "code" in payload and payload["code"] != function.code:
        if db.scalar(select(SystemFunction).where(SystemFunction.code == payload["code"])):
            raise HTTPException(status_code=409, detail="Function code already exists")
    for field in ("name", "code", "module", "description", "is_active"):
        if field in payload:
            setattr(function, field, payload[field])
    function.updated_at = datetime.utcnow()
    db.add(function)
    db.commit()
    db.refresh(function)
    return function


def delete_system_function(db: Session, function_id: int) -> None:
    function = _get_function_or_404(db, function_id)
    db.delete(function)
    db.commit()


def list_issuing_units(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    is_active: bool | None,
) -> tuple[list[IssuingUnit], int]:
    statement: Select[tuple[IssuingUnit]] = select(IssuingUnit).options(selectinload(IssuingUnit.parent))
    count_statement = select(func.count()).select_from(IssuingUnit)
    if search:
        criteria = or_(
            IssuingUnit.code.ilike(f"%{search}%"),
            IssuingUnit.name.ilike(f"%{search}%"),
            IssuingUnit.short_name.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if is_active is not None:
        statement = statement.where(IssuingUnit.is_active == is_active)
        count_statement = count_statement.where(IssuingUnit.is_active == is_active)
    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(IssuingUnit.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_issuing_unit(db: Session, payload: dict[str, Any]) -> IssuingUnit:
    if db.scalar(select(IssuingUnit).where(IssuingUnit.code == payload["code"])):
        raise HTTPException(status_code=409, detail="Issuing unit code already exists")
    if payload.get("parent_id") is not None:
        _get_issuing_unit_or_404(db, payload["parent_id"])
    item = IssuingUnit(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def update_issuing_unit(db: Session, unit_id: int, payload: dict[str, Any]) -> IssuingUnit:
    item = _get_issuing_unit_or_404(db, unit_id)
    if "code" in payload and payload["code"] != item.code:
        if db.scalar(select(IssuingUnit).where(IssuingUnit.code == payload["code"])):
            raise HTTPException(status_code=409, detail="Issuing unit code already exists")
    if "parent_id" in payload:
        if payload["parent_id"] == unit_id:
            raise HTTPException(status_code=400, detail="Parent cannot be self")
        if payload["parent_id"] is not None:
            _get_issuing_unit_or_404(db, payload["parent_id"])
    for field, value in payload.items():
        setattr(item, field, value)
    item.updated_at = datetime.utcnow()
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def delete_issuing_unit(db: Session, unit_id: int) -> None:
    item = _get_issuing_unit_or_404(db, unit_id)
    if item.children or item.departments or item.work_assignment_documents or item.notice_documents:
        raise HTTPException(status_code=400, detail="Issuing unit is in use")
    db.delete(item)
    db.commit()


def list_departments(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    issuing_unit_id: int | None,
    is_active: bool | None,
) -> tuple[list[Department], int]:
    statement: Select[tuple[Department]] = select(Department).options(
        selectinload(Department.issuing_unit),
        selectinload(Department.parent),
    )
    count_statement = select(func.count()).select_from(Department)
    if search:
        criteria = or_(
            Department.code.ilike(f"%{search}%"),
            Department.name.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if issuing_unit_id is not None:
        statement = statement.where(Department.issuing_unit_id == issuing_unit_id)
        count_statement = count_statement.where(Department.issuing_unit_id == issuing_unit_id)
    if is_active is not None:
        statement = statement.where(Department.is_active == is_active)
        count_statement = count_statement.where(Department.is_active == is_active)
    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(Department.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_department(db: Session, payload: dict[str, Any]) -> Department:
    if db.scalar(select(Department).where(Department.code == payload["code"])):
        raise HTTPException(status_code=409, detail="Department code already exists")
    if payload.get("issuing_unit_id") is not None:
        _get_issuing_unit_or_404(db, payload["issuing_unit_id"])
    if payload.get("parent_id") is not None:
        _get_department_or_404(db, payload["parent_id"])
    item = Department(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def update_department(db: Session, department_id: int, payload: dict[str, Any]) -> Department:
    item = _get_department_or_404(db, department_id)
    if "code" in payload and payload["code"] != item.code:
        if db.scalar(select(Department).where(Department.code == payload["code"])):
            raise HTTPException(status_code=409, detail="Department code already exists")
    if "issuing_unit_id" in payload and payload["issuing_unit_id"] is not None:
        _get_issuing_unit_or_404(db, payload["issuing_unit_id"])
    if "parent_id" in payload:
        if payload["parent_id"] == department_id:
            raise HTTPException(status_code=400, detail="Parent cannot be self")
        if payload["parent_id"] is not None:
            _get_department_or_404(db, payload["parent_id"])
    for field, value in payload.items():
        setattr(item, field, value)
    item.updated_at = datetime.utcnow()
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def delete_department(db: Session, department_id: int) -> None:
    item = _get_department_or_404(db, department_id)
    if item.children or item.positions or item.users or item.work_items or item.work_assignment_documents or item.assigned_work_documents or item.notice_documents:
        raise HTTPException(status_code=400, detail="Department is in use")
    db.delete(item)
    db.commit()


def list_positions(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    department_id: int | None,
    is_active: bool | None,
) -> tuple[list[Position], int]:
    statement: Select[tuple[Position]] = select(Position).options(selectinload(Position.department))
    count_statement = select(func.count()).select_from(Position)
    if search:
        criteria = or_(
            Position.code.ilike(f"%{search}%"),
            Position.name.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if department_id is not None:
        statement = statement.where(Position.department_id == department_id)
        count_statement = count_statement.where(Position.department_id == department_id)
    if is_active is not None:
        statement = statement.where(Position.is_active == is_active)
        count_statement = count_statement.where(Position.is_active == is_active)
    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(Position.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_position(db: Session, payload: dict[str, Any]) -> Position:
    if db.scalar(select(Position).where(Position.code == payload["code"])):
        raise HTTPException(status_code=409, detail="Position code already exists")
    if payload.get("department_id") is not None:
        _get_department_or_404(db, payload["department_id"])
    item = Position(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def update_position(db: Session, position_id: int, payload: dict[str, Any]) -> Position:
    item = _get_position_or_404(db, position_id)
    if "code" in payload and payload["code"] != item.code:
        if db.scalar(select(Position).where(Position.code == payload["code"])):
            raise HTTPException(status_code=409, detail="Position code already exists")
    if "department_id" in payload and payload["department_id"] is not None:
        _get_department_or_404(db, payload["department_id"])
    for field, value in payload.items():
        setattr(item, field, value)
    item.updated_at = datetime.utcnow()
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def delete_position(db: Session, position_id: int) -> None:
    item = _get_position_or_404(db, position_id)
    if item.users or item.work_items:
        raise HTTPException(status_code=400, detail="Position is in use")
    db.delete(item)
    db.commit()


def list_work_documents(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    status_value: str | None,
    issuing_unit_id: int | None,
    department_id: int | None,
    assigned_department_id: int | None,
) -> tuple[list[WorkAssignmentDocument], int]:
    statement: Select[tuple[WorkAssignmentDocument]] = select(WorkAssignmentDocument).options(
        selectinload(WorkAssignmentDocument.issuing_unit),
        selectinload(WorkAssignmentDocument.department),
        selectinload(WorkAssignmentDocument.assigned_department),
        selectinload(WorkAssignmentDocument.assigned_by_user),
        selectinload(WorkAssignmentDocument.work_items),
    )
    count_statement = select(func.count()).select_from(WorkAssignmentDocument)
    if search:
        criteria = or_(
            WorkAssignmentDocument.document_code.ilike(f"%{search}%"),
            WorkAssignmentDocument.title.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if status_value:
        statement = statement.where(WorkAssignmentDocument.status == status_value)
        count_statement = count_statement.where(WorkAssignmentDocument.status == status_value)
    if issuing_unit_id is not None:
        statement = statement.where(WorkAssignmentDocument.issuing_unit_id == issuing_unit_id)
        count_statement = count_statement.where(WorkAssignmentDocument.issuing_unit_id == issuing_unit_id)
    if department_id is not None:
        statement = statement.where(WorkAssignmentDocument.department_id == department_id)
        count_statement = count_statement.where(WorkAssignmentDocument.department_id == department_id)
    if assigned_department_id is not None:
        statement = statement.where(WorkAssignmentDocument.assigned_department_id == assigned_department_id)
        count_statement = count_statement.where(WorkAssignmentDocument.assigned_department_id == assigned_department_id)
    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(WorkAssignmentDocument.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_work_document(db: Session, payload: dict[str, Any]) -> WorkAssignmentDocument:
    if db.scalar(select(WorkAssignmentDocument).where(WorkAssignmentDocument.document_code == payload["document_code"])):
        raise HTTPException(status_code=409, detail="Work document code already exists")
    _validate_work_document_relations(db, payload)
    item = WorkAssignmentDocument(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def update_work_document(db: Session, document_id: int, payload: dict[str, Any]) -> WorkAssignmentDocument:
    item = _get_work_document_or_404(db, document_id)
    if "document_code" in payload and payload["document_code"] != item.document_code:
        if db.scalar(select(WorkAssignmentDocument).where(WorkAssignmentDocument.document_code == payload["document_code"])):
            raise HTTPException(status_code=409, detail="Work document code already exists")
    _validate_work_document_relations(db, payload)
    for field, value in payload.items():
        setattr(item, field, value)
    item.updated_at = datetime.utcnow()
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def delete_work_document(db: Session, document_id: int) -> None:
    item = _get_work_document_or_404(db, document_id)
    db.delete(item)
    db.commit()


def list_work_items(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    status_value: str | None,
    work_document_id: int | None,
    assignee_user_id: int | None,
    department_id: int | None,
) -> tuple[list[WorkItem], int]:
    statement: Select[tuple[WorkItem]] = select(WorkItem).options(
        selectinload(WorkItem.work_document),
        selectinload(WorkItem.assignee),
        selectinload(WorkItem.department),
        selectinload(WorkItem.position),
    )
    count_statement = select(func.count()).select_from(WorkItem)
    if search:
        criteria = or_(
            WorkItem.title.ilike(f"%{search}%"),
            WorkItem.description.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if status_value:
        statement = statement.where(WorkItem.status == status_value)
        count_statement = count_statement.where(WorkItem.status == status_value)
    if work_document_id is not None:
        statement = statement.where(WorkItem.work_document_id == work_document_id)
        count_statement = count_statement.where(WorkItem.work_document_id == work_document_id)
    if assignee_user_id is not None:
        statement = statement.where(WorkItem.assignee_user_id == assignee_user_id)
        count_statement = count_statement.where(WorkItem.assignee_user_id == assignee_user_id)
    if department_id is not None:
        statement = statement.where(WorkItem.department_id == department_id)
        count_statement = count_statement.where(WorkItem.department_id == department_id)
    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(WorkItem.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_work_item(db: Session, payload: dict[str, Any]) -> WorkItem:
    _validate_work_item_relations(db, payload)
    item = WorkItem(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def update_work_item(db: Session, item_id: int, payload: dict[str, Any]) -> WorkItem:
    item = _get_work_item_or_404(db, item_id)
    _validate_work_item_relations(db, payload)
    for field, value in payload.items():
        setattr(item, field, value)
    item.updated_at = datetime.utcnow()
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def delete_work_item(db: Session, item_id: int) -> None:
    item = _get_work_item_or_404(db, item_id)
    db.delete(item)
    db.commit()


def list_notice_documents(
    db: Session,
    page: int,
    page_size: int,
    search: str | None,
    status_value: str | None,
    issuing_unit_id: int | None,
    department_id: int | None,
) -> tuple[list[NoticeDocument], int]:
    statement: Select[tuple[NoticeDocument]] = select(NoticeDocument).options(
        selectinload(NoticeDocument.issuing_unit),
        selectinload(NoticeDocument.department),
        selectinload(NoticeDocument.posted_by_user),
    )
    count_statement = select(func.count()).select_from(NoticeDocument)
    if search:
        criteria = or_(
            NoticeDocument.notice_code.ilike(f"%{search}%"),
            NoticeDocument.title.ilike(f"%{search}%"),
        )
        statement = statement.where(criteria)
        count_statement = count_statement.where(criteria)
    if status_value:
        statement = statement.where(NoticeDocument.status == status_value)
        count_statement = count_statement.where(NoticeDocument.status == status_value)
    if issuing_unit_id is not None:
        statement = statement.where(NoticeDocument.issuing_unit_id == issuing_unit_id)
        count_statement = count_statement.where(NoticeDocument.issuing_unit_id == issuing_unit_id)
    if department_id is not None:
        statement = statement.where(NoticeDocument.department_id == department_id)
        count_statement = count_statement.where(NoticeDocument.department_id == department_id)
    total = db.scalar(count_statement) or 0
    items = list(db.scalars(
        statement.order_by(NoticeDocument.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total


def create_notice_document(db: Session, payload: dict[str, Any]) -> NoticeDocument:
    if db.scalar(select(NoticeDocument).where(NoticeDocument.notice_code == payload["notice_code"])):
        raise HTTPException(status_code=409, detail="Notice code already exists")
    _validate_notice_relations(db, payload)
    item = NoticeDocument(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def update_notice_document(db: Session, notice_id: int, payload: dict[str, Any]) -> NoticeDocument:
    item = _get_notice_document_or_404(db, notice_id)
    if "notice_code" in payload and payload["notice_code"] != item.notice_code:
        if db.scalar(select(NoticeDocument).where(NoticeDocument.notice_code == payload["notice_code"])):
            raise HTTPException(status_code=409, detail="Notice code already exists")
    _validate_notice_relations(db, payload)
    for field, value in payload.items():
        setattr(item, field, value)
    item.updated_at = datetime.utcnow()
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def delete_notice_document(db: Session, notice_id: int) -> None:
    item = _get_notice_document_or_404(db, notice_id)
    db.delete(item)
    db.commit()


def create_backup(db: Session, actor_user_id: int) -> dict[str, Any]:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_root = Path(settings.BACKUP_DIR) / timestamp
    backup_root.mkdir(parents=True, exist_ok=True)

    snapshot = _build_db_snapshot(db)
    snapshot_path = backup_root / "db_snapshot.json"
    snapshot_path.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    copied_items = [snapshot_path]
    db_dump_path = Path("db_dump.sql")
    if db_dump_path.exists():
        destination = backup_root / db_dump_path.name
        shutil.copy2(db_dump_path, destination)
        copied_items.append(destination)

    faiss_path = Path(settings.FAISS_INDEX_PATH)
    if faiss_path.exists():
        destination = backup_root / faiss_path.name
        shutil.copytree(faiss_path, destination, dirs_exist_ok=True)
        copied_items.append(destination)

    metadata = {
        "backup_name": timestamp,
        "backup_path": str(backup_root),
        "created_at": datetime.now().isoformat(),
        "created_by": actor_user_id,
        "has_db_snapshot": True,
        "has_db_dump": db_dump_path.exists(),
        "has_faiss_index": faiss_path.exists(),
    }
    (backup_root / "backup_meta.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    size_kb = round(sum(_path_size(item) for item in copied_items) / 1024, 2)
    write_system_log(
        db,
        action="manual_backup",
        status_code=200,
        user_id=actor_user_id,
        detail=json.dumps(metadata, ensure_ascii=False),
    )
    return {
        "backup_name": timestamp,
        "backup_path": str(backup_root),
        "created_at": metadata["created_at"],
        "size_kb": size_kb,
        "has_db_dump": metadata["has_db_dump"],
        "has_faiss_index": metadata["has_faiss_index"],
    }


def list_backups() -> list[dict[str, Any]]:
    backup_dir = Path(settings.BACKUP_DIR)
    backup_dir.mkdir(parents=True, exist_ok=True)
    items: list[dict[str, Any]] = []
    for entry in sorted(backup_dir.iterdir(), reverse=True):
        if not entry.is_dir():
            continue
        meta_path = entry / "backup_meta.json"
        if meta_path.exists():
            items.append(json.loads(meta_path.read_text(encoding="utf-8")))
        else:
            items.append(
                {
                    "backup_name": entry.name,
                    "backup_path": str(entry),
                    "created_at": datetime.fromtimestamp(entry.stat().st_mtime).isoformat(),
                    "has_db_snapshot": (entry / "db_snapshot.json").exists(),
                    "has_db_dump": (entry / "db_dump.sql").exists(),
                    "has_faiss_index": (entry / Path(settings.FAISS_INDEX_PATH).name).exists(),
                }
            )
    return items


def restore_backup(
    db: Session,
    backup_name: str,
    actor_user_id: int,
) -> dict[str, Any]:
    backup_root = Path(settings.BACKUP_DIR) / backup_name
    if not backup_root.exists():
        raise HTTPException(status_code=404, detail="Backup not found")

    snapshot_path = backup_root / "db_snapshot.json"
    if not snapshot_path.exists():
        raise HTTPException(status_code=404, detail="Backup snapshot not found")

    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    _restore_db_snapshot(db, snapshot)

    backup_faiss = backup_root / Path(settings.FAISS_INDEX_PATH).name
    if backup_faiss.exists():
        target = Path(settings.FAISS_INDEX_PATH)
        if target.exists():
            shutil.rmtree(target)
        shutil.copytree(backup_faiss, target)

    backup_dump = backup_root / "db_dump.sql"
    if backup_dump.exists():
        shutil.copy2(backup_dump, Path("db_dump.sql"))

    detail = {
        "backup_name": backup_name,
        "restored_at": datetime.now().isoformat(),
    }
    write_system_log(
        db,
        action="restore_backup",
        status_code=200,
        user_id=actor_user_id,
        detail=json.dumps(detail, ensure_ascii=False),
    )
    return detail


def write_system_log(
    db: Session,
    action: str,
    status_code: int,
    user_id: int | None = None,
    detail: str | None = None,
    ip_address: str | None = None,
) -> None:
    db.add(
        SystemLog(
            user_id=user_id,
            action=action,
            detail=detail,
            ip_address=ip_address,
            status_code=status_code,
        )
    )
    db.commit()


def _assert_unique_user(db: Session, username: str, email: str | None) -> None:
    if db.scalar(select(User).where(User.username == username)):
        raise HTTPException(status_code=409, detail="Username already exists")
    if email and db.scalar(select(User).where(User.email == email)):
        raise HTTPException(status_code=409, detail="Email already exists")


def _validate_user_relations(
    db: Session,
    permission_group_id: int | None,
    department_id: int | None,
    position_id: int | None,
) -> None:
    if permission_group_id is not None:
        _get_group_or_404(db, permission_group_id)
    if department_id is not None:
        _get_department_or_404(db, department_id)
    if position_id is not None:
        position = _get_position_or_404(db, position_id)
        if department_id is not None and position.department_id not in {None, department_id}:
            raise HTTPException(status_code=400, detail="Position does not belong to selected department")


def _get_user_or_404(db: Session, user_id: int) -> User:
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user


def _get_config_or_404(db: Session, config_id: int) -> SystemConfig:
    config = db.get(SystemConfig, config_id)
    if not config:
        raise HTTPException(status_code=404, detail="Config not found")
    return config


def _get_api_key_or_404(db: Session, api_key_id: int) -> APIKey:
    api_key = db.get(APIKey, api_key_id)
    if not api_key:
        raise HTTPException(status_code=404, detail="API key not found")
    return api_key


def _get_group_or_404(db: Session, group_id: int) -> PermissionGroup:
    group = db.get(PermissionGroup, group_id)
    if not group:
        raise HTTPException(status_code=404, detail="Permission group not found")
    return group


def _get_function_or_404(db: Session, function_id: int) -> SystemFunction:
    function = db.get(SystemFunction, function_id)
    if not function:
        raise HTTPException(status_code=404, detail="System function not found")
    return function


def _get_issuing_unit_or_404(db: Session, unit_id: int) -> IssuingUnit:
    item = db.get(IssuingUnit, unit_id)
    if not item:
        raise HTTPException(status_code=404, detail="Issuing unit not found")
    return item


def _get_department_or_404(db: Session, department_id: int) -> Department:
    item = db.get(Department, department_id)
    if not item:
        raise HTTPException(status_code=404, detail="Department not found")
    return item


def _get_position_or_404(db: Session, position_id: int) -> Position:
    item = db.get(Position, position_id)
    if not item:
        raise HTTPException(status_code=404, detail="Position not found")
    return item


def _get_work_document_or_404(db: Session, document_id: int) -> WorkAssignmentDocument:
    item = db.get(WorkAssignmentDocument, document_id)
    if not item:
        raise HTTPException(status_code=404, detail="Work document not found")
    return item


def _get_work_item_or_404(db: Session, item_id: int) -> WorkItem:
    item = db.get(WorkItem, item_id)
    if not item:
        raise HTTPException(status_code=404, detail="Work item not found")
    return item


def _get_notice_document_or_404(db: Session, notice_id: int) -> NoticeDocument:
    item = db.get(NoticeDocument, notice_id)
    if not item:
        raise HTTPException(status_code=404, detail="Notice document not found")
    return item


def _replace_group_permissions(
    db: Session,
    group: PermissionGroup,
    permissions: list[dict[str, Any]],
) -> None:
    existing_permissions = db.scalars(
        select(PermissionGroupFunction).where(PermissionGroupFunction.group_id == group.id)
    ).all()
    for item in existing_permissions:
        db.delete(item)
    db.flush()

    for permission in permissions:
        function = db.get(SystemFunction, permission["function_id"])
        if not function:
            raise HTTPException(status_code=404, detail="System function not found")
        db.add(
            PermissionGroupFunction(
                group_id=group.id,
                function_id=permission["function_id"],
                can_view=permission.get("can_view", True),
                can_create=permission.get("can_create", False),
                can_update=permission.get("can_update", False),
                can_delete=permission.get("can_delete", False),
            )
        )


def _validate_work_document_relations(db: Session, payload: dict[str, Any]) -> None:
    if payload.get("issuing_unit_id") is not None:
        _get_issuing_unit_or_404(db, payload["issuing_unit_id"])
    if payload.get("department_id") is not None:
        _get_department_or_404(db, payload["department_id"])
    if payload.get("assigned_department_id") is not None:
        _get_department_or_404(db, payload["assigned_department_id"])
    if payload.get("assigned_by_user_id") is not None:
        _get_user_or_404(db, payload["assigned_by_user_id"])


def _validate_work_item_relations(db: Session, payload: dict[str, Any]) -> None:
    if payload.get("work_document_id") is not None:
        _get_work_document_or_404(db, payload["work_document_id"])
    if payload.get("assignee_user_id") is not None:
        _get_user_or_404(db, payload["assignee_user_id"])
    if payload.get("department_id") is not None:
        _get_department_or_404(db, payload["department_id"])
    if payload.get("position_id") is not None:
        _get_position_or_404(db, payload["position_id"])


def _validate_notice_relations(db: Session, payload: dict[str, Any]) -> None:
    if payload.get("issuing_unit_id") is not None:
        _get_issuing_unit_or_404(db, payload["issuing_unit_id"])
    if payload.get("department_id") is not None:
        _get_department_or_404(db, payload["department_id"])
    if payload.get("posted_by_user_id") is not None:
        _get_user_or_404(db, payload["posted_by_user_id"])


def _sync_config_to_settings(config_key: str, config_value: str, data_type: str) -> None:
    mapping = {
        "MODEL_NAME": "MODEL_NAME",
        "MODEL_PATH": "MODEL_PATH",
        "FAISS_INDEX_PATH": "FAISS_INDEX_PATH",
        "VECTOR_DIM": "VECTOR_DIM",
        "CHUNK_SIZE": "CHUNK_SIZE",
        "CHUNK_OVERLAP": "CHUNK_OVERLAP",
        "OCR_LANG": "OCR_LANG",
        "UPLOAD_DIR": "UPLOAD_DIR",
        "MAX_FILE_SIZE_MB": "MAX_FILE_SIZE_MB",
        "BACKUP_DIR": "BACKUP_DIR",
        "SSO_ENABLED": "SSO_ENABLED",
        "SSO_PROVIDER_NAME": "SSO_PROVIDER_NAME",
        "SSO_SHARED_SECRET": "SSO_SHARED_SECRET",
        "SSO_AUTO_CREATE_USERS": "SSO_AUTO_CREATE_USERS",
    }
    setting_name = mapping.get(config_key.upper())
    if not setting_name:
        return

    value: Any = config_value
    normalized_type = data_type.lower()
    if normalized_type == "int":
        value = int(config_value)
    elif normalized_type == "bool":
        value = config_value.lower() in {"1", "true", "yes", "on"}
    setattr(settings, setting_name, value)


def _build_db_snapshot(db: Session) -> dict[str, list[dict[str, Any]]]:
    snapshot: dict[str, list[dict[str, Any]]] = {}
    for model in SNAPSHOT_MODELS:
        rows = db.scalars(select(model)).all()
        snapshot[model.__tablename__] = [_serialize_model_row(model, row) for row in rows]
    return snapshot


def _restore_db_snapshot(db: Session, snapshot: dict[str, list[dict[str, Any]]]) -> None:
    for model in SNAPSHOT_DELETE_ORDER:
        db.execute(delete(model))
    db.commit()

    insert_order = SNAPSHOT_MODELS
    for model in insert_order:
        rows = snapshot.get(model.__tablename__, [])
        if not rows:
            continue
        restored_rows = [_deserialize_model_row(model, row) for row in rows]
        db.execute(model.__table__.insert(), restored_rows)
    db.commit()


def _serialize_model_row(model: type, row: Any) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for column in model.__table__.columns:
        value = getattr(row, column.name)
        if isinstance(value, (datetime, date)):
            data[column.name] = value.isoformat()
        else:
            data[column.name] = value
    return data


def _deserialize_model_row(model: type, row: dict[str, Any]) -> dict[str, Any]:
    restored: dict[str, Any] = {}
    for column in model.__table__.columns:
        value = row.get(column.name)
        if value is None:
            restored[column.name] = None
            continue
        try:
            python_type = column.type.python_type
        except (AttributeError, NotImplementedError):
            python_type = None
        if python_type is datetime and isinstance(value, str):
            restored[column.name] = datetime.fromisoformat(value)
        elif python_type is date and isinstance(value, str):
            restored[column.name] = date.fromisoformat(value)
        else:
            restored[column.name] = value
    return restored


def _count_files(path: Path) -> int:
    if not path.exists():
        return 0
    return sum(1 for item in path.iterdir() if item.is_file())


def _path_size(path: Path) -> int:
    if path.is_file():
        return path.stat().st_size
    return sum(item.stat().st_size for item in path.rglob("*") if item.is_file())


def _get_workbook_class():
    from openpyxl import Workbook

    return Workbook
